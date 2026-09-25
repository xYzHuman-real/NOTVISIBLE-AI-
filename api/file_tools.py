import csv
import io
from pathlib import Path

def extract_text(filename: str, data: bytes) -> str:
    suffix = Path(filename).suffix.lower()
    if suffix in {".txt", ".md", ".csv", ".json", ".py", ".js", ".ts", ".html", ".css"}:
        return data.decode("utf-8", errors="replace")
    if suffix == ".pdf":
        from pypdf import PdfReader
        reader = PdfReader(io.BytesIO(data))
        return "\n\n".join(page.extract_text() or "" for page in reader.pages)
    if suffix == ".docx":
        from docx import Document
        doc = Document(io.BytesIO(data))
        return "\n".join(p.text for p in doc.paragraphs)
    if suffix == ".xlsx":
        from openpyxl import load_workbook
        wb = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
        lines = []
        for ws in wb.worksheets:
            lines.append(f"[Sheet: {ws.title}]")
            for row in ws.iter_rows(values_only=True):
                lines.append("\t".join("" if v is None else str(v) for v in row))
        return "\n".join(lines)
    raise ValueError("Unsupported file type")
